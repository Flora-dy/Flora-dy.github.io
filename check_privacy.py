# -*- coding: utf-8 -*-
"""发布前自检：确认要上公网的文件里没有真实人员信息。

这个站是公开的，而工具包目录里就放着真实人员表，复制时很容易带出去。
所以每次发布前扫一遍：把本机数据表里的所有姓名当作黑名单，
在待发布文件里逐个搜；命中就中止发布。
"""
import os, re, sys, glob, zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
# 真实人员数据在哪儿找（只读、不上传）——按需增删
DATA_DIRS = ["/Users/zzp/Downloads/应用技术中心架构"]
PUBLISH = ["index.html", "README.md", "使用说明.md", "人员信息填写模板.xlsx"]
# 这些文件本身就是示例/模板，里面的名字不是真人，不能拿来当黑名单
SKIP_SRC = ("人员信息填写模板.xlsx", "组织列填写对照.xlsx")
# 教科书式化名，出现在哪里都不算泄露
PLACEHOLDER = {"张三", "李四", "王五", "赵六", "钱七", "孙八", "周九", "吴十",
               "郑一", "王二", "冯佳", "陈五"}

def real_names():
    """把本机所有人员表里的「姓名」列收集成黑名单"""
    try:
        import openpyxl
    except ImportError:
        print("  （未装 openpyxl，跳过姓名比对）")
        return set()
    names = set()
    for d in DATA_DIRS:
        for p in glob.glob(os.path.join(d, "**", "*.xlsx"), recursive=True):
            _b = os.path.basename(p)
            if _b.startswith("~$") or _b in SKIP_SRC:
                continue
            try:
                wb = openpyxl.load_workbook(p, data_only=True, read_only=True)
            except Exception:
                continue
            for ws in wb.worksheets:
                rows = ws.iter_rows(max_row=1, values_only=True)
                hdr = next(rows, None) or ()
                hdr = [str(v).strip() if v else "" for v in hdr]
                if "姓名" not in hdr:
                    continue
                j = hdr.index("姓名")
                for r in ws.iter_rows(min_row=2, values_only=True):
                    if j < len(r) and r[j]:
                        v = str(r[j]).strip()
                        # 只收 2~4 字的中文姓名，避免「待招管培生1」这类误判
                        if 2 <= len(v) <= 4 and re.fullmatch(r"[一-龥]+", v) \
                           and v not in PLACEHOLDER:
                            names.add(v)
            wb.close()
    return names

def text_of(path):
    if path.endswith(".xlsx"):
        out = []
        with zipfile.ZipFile(path) as z:
            for n in z.namelist():
                if n.endswith(".xml"):
                    out.append(z.read(n).decode("utf8", "ignore"))
        return "".join(out)
    with open(path, encoding="utf8", errors="ignore") as f:
        return f.read()

def main():
    names = real_names()
    print(f"  黑名单 {len(names)} 个姓名，检查 {len(PUBLISH)} 个待发布文件")
    bad = []
    for f in PUBLISH:
        p = os.path.join(HERE, f)
        if not os.path.exists(p):
            print(f"  ⚠️ 缺文件：{f}")
            continue
        s = text_of(p)
        hit = sorted(n for n in names if n in s)
        if hit:
            bad.append((f, hit))
    # 仓库里不该出现的文件类型
    for p in glob.glob(os.path.join(HERE, "*")):
        b = os.path.basename(p)
        if b.endswith((".docx", ".pptx", ".pdf")) or \
           (b.endswith(".xlsx") and b != "人员信息填写模板.xlsx"):
            bad.append((b, ["这类文件不该进公开仓库"]))
    if bad:
        print("\n  ❌ 发现可能的真实信息：")
        for f, hit in bad:
            print(f"     {f}: {'、'.join(hit[:10])}" + ("…" if len(hit) > 10 else ""))
        return 1
    print("  ✅ 未发现真实人员信息")
    return 0

if __name__ == "__main__":
    sys.exit(main())
